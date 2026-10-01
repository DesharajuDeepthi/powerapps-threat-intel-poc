from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.excel_writer import append_threat_intel_records
from src.message_parser import ThreatIntelRecord


def test_append_threat_intel_records_creates_workbook(tmp_path: Path) -> None:
    output_file = tmp_path / "Threat_Intelligence_Findings.xlsx"
    records = [
        ThreatIntelRecord(
            slack_message_ts="1.000",
            document="first.pdf",
            threat="Hive Ransomware",
            impacted_technology="FortiOS",
            cve_count=1,
            cves=["CVE-2024-1111"],
        ),
        ThreatIntelRecord(
            slack_message_ts="2.000",
            document="second.pdf",
            threat="Example Threat",
            impacted_technology="Microsoft Exchange",
            cve_count=2,
            cves=["CVE-2024-2222", "CVE-2024-3333"],
        ),
    ]

    result = append_threat_intel_records(records, output_file)

    assert result.rows_added == 2
    workbook = load_workbook(output_file)
    worksheet = workbook["Slack Messages"]
    assert worksheet.max_row == 3
    assert worksheet["B2"].value == "1.000"
    assert worksheet["C2"].value == "first.pdf"
    assert worksheet["G3"].value == "CVE-2024-2222; CVE-2024-3333"


def test_append_threat_intel_records_skips_duplicate_message_ts(tmp_path: Path) -> None:
    output_file = tmp_path / "Threat_Intelligence_Findings.xlsx"
    records = [
        ThreatIntelRecord(
            slack_message_ts="1.000",
            document="first.pdf",
            threat="Hive Ransomware",
            impacted_technology="FortiOS",
            cve_count=1,
            cves=["CVE-2024-1111"],
        )
    ]

    first_result = append_threat_intel_records(records, output_file)
    second_result = append_threat_intel_records(records, output_file)

    assert first_result.rows_added == 1
    assert second_result.rows_added == 0
    assert second_result.rows_skipped == 1
    workbook = load_workbook(output_file)
    assert workbook["Slack Messages"].max_row == 2
