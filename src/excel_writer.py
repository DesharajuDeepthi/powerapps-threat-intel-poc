from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from src.models import DocumentAnalysis
from src.message_parser import ThreatIntelRecord

LOGGER = logging.getLogger(__name__)

WORKSHEET_NAME = "Slack Messages"
DOCUMENT_ANALYSIS_WORKSHEET_NAME = "Document Analysis"
HEADERS = [
    "Processing_Date",
    "Slack_Message_TS",
    "Document",
    "Threat",
    "Impacted_Technology",
    "CVE_Count",
    "CVE",
]
DOCUMENT_ANALYSIS_HEADERS = [
    "Processing_Date",
    "Slack_Message_Date",
    "Document_Name",
    "Document_ID",
    "Threat_Name",
    "CVE",
    "Technology",
    "Vendor",
    "Semantic_Category",
    "Document_Relevance",
    "Active_Exploitation_Score",
    "Initial_Access_Score",
    "RCE_Score",
    "Authentication_Bypass_Score",
    "Background_Reference_Score",
    "IP_Addresses",
    "Domains",
    "URLs",
    "Hashes",
    "MITRE_Techniques",
    "Container_Indicators",
    "Evidence_Text",
    "Analyst_Comment",
    "PDF_Link",
    "Slack_Message_ID",
    "Slack_File_ID",
]


@dataclass(frozen=True)
class ExcelWriteResult:
    output_file: Path
    records_seen: int
    rows_added: int
    rows_skipped: int


def append_threat_intel_records(
    records: list[ThreatIntelRecord],
    output_file: Path,
) -> ExcelWriteResult:
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(output_file) if output_file.exists() else Workbook()
    worksheet = _worksheet(workbook)
    existing_message_ts = _existing_message_timestamps(worksheet)

    rows_added = 0
    rows_skipped = 0
    processing_date = datetime.now(timezone.utc).date().isoformat()

    for record in records:
        if record.slack_message_ts in existing_message_ts:
            rows_skipped += 1
            continue

        worksheet.append(
            [
                processing_date,
                record.slack_message_ts,
                record.document,
                record.threat,
                record.impacted_technology,
                record.cve_count,
                "; ".join(record.cves),
            ]
        )
        existing_message_ts.add(record.slack_message_ts)
        rows_added += 1

    workbook.save(output_file)
    LOGGER.info(
        "Excel workbook updated at %s: %s rows added, %s skipped",
        output_file,
        rows_added,
        rows_skipped,
    )

    return ExcelWriteResult(
        output_file=output_file,
        records_seen=len(records),
        rows_added=rows_added,
        rows_skipped=rows_skipped,
    )


def refresh_document_analysis_sheet(
    analyses: list[DocumentAnalysis],
    output_file: Path,
) -> None:
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(output_file) if output_file.exists() else Workbook()
    if DOCUMENT_ANALYSIS_WORKSHEET_NAME in workbook.sheetnames:
        del workbook[DOCUMENT_ANALYSIS_WORKSHEET_NAME]

    worksheet = workbook.create_sheet(DOCUMENT_ANALYSIS_WORKSHEET_NAME)
    for index, header in enumerate(DOCUMENT_ANALYSIS_HEADERS, start=1):
        worksheet.cell(row=1, column=index, value=header)

    for analysis in sorted(analyses, key=lambda item: item.processed_at):
        for cve in analysis.cves:
            row = [
                _date_part(analysis.processed_at),
                _slack_ts_to_iso(analysis.slack_message_timestamp),
                analysis.document_name,
                analysis.document_id,
                analysis.threat_name or "",
                cve.cve,
                cve.technology,
                cve.vendor,
                cve.semantic_category,
                cve.document_relevance,
                cve.semantic_scores.get("active_exploitation", 0.0),
                cve.semantic_scores.get("initial_access", 0.0),
                cve.semantic_scores.get("remote_code_execution", 0.0),
                cve.semantic_scores.get("authentication_bypass", 0.0),
                cve.semantic_scores.get("background_reference", 0.0),
                "; ".join(analysis.indicators.ipv4),
                "; ".join(analysis.indicators.domains),
                "; ".join(analysis.indicators.urls),
                _hashes(analysis),
                "; ".join(analysis.indicators.mitre_techniques),
                "; ".join(analysis.indicators.container_indicators),
                cve.evidence_text,
                cve.analyst_comment,
                "Open PDF",
                analysis.slack_message_id or "",
                analysis.slack_file_id or "",
            ]
            worksheet.append(row)
            pdf_cell = worksheet.cell(row=worksheet.max_row, column=24)
            pdf_cell.hyperlink = _relative_pdf_link(output_file, analysis.local_file_path)
            pdf_cell.font = Font(color="0000EE", underline="single")

    workbook.save(output_file)
    LOGGER.info("Document analysis Excel worksheet refreshed at %s", output_file)


def _worksheet(workbook: Workbook) -> Worksheet:
    if WORKSHEET_NAME in workbook.sheetnames:
        worksheet = workbook[WORKSHEET_NAME]
        _ensure_headers(worksheet)
        return worksheet

    if workbook.active.max_row == 1 and workbook.active.max_column == 1 and workbook.active["A1"].value is None:
        worksheet = workbook.active
        worksheet.title = WORKSHEET_NAME
    else:
        worksheet = workbook.create_sheet(WORKSHEET_NAME)

    _write_headers(worksheet)
    return worksheet


def _ensure_headers(worksheet: Worksheet) -> None:
    current_headers = [worksheet.cell(row=1, column=index).value for index in range(1, len(HEADERS) + 1)]
    if current_headers == HEADERS:
        return

    if worksheet.max_row == 1 and all(value is None for value in current_headers):
        _write_headers(worksheet)


def _write_headers(worksheet: Worksheet) -> None:
    for index, header in enumerate(HEADERS, start=1):
        worksheet.cell(row=1, column=index, value=header)


def _existing_message_timestamps(worksheet: Worksheet) -> set[str]:
    headers = [worksheet.cell(row=1, column=index).value for index in range(1, worksheet.max_column + 1)]
    try:
        timestamp_column = headers.index("Slack_Message_TS") + 1
    except ValueError:
        return set()

    values: set[str] = set()
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=timestamp_column).value
        if value is not None:
            values.add(str(value))
    return values


def _date_part(value: str) -> str:
    return value.split("T", 1)[0]


def _slack_ts_to_iso(value: str | None) -> str:
    if not value:
        return ""
    try:
        timestamp = float(value)
    except ValueError:
        return value
    return datetime.fromtimestamp(timestamp, tz=timezone.utc).isoformat()


def _hashes(analysis: DocumentAnalysis) -> str:
    return "; ".join([*analysis.indicators.sha256, *analysis.indicators.sha1, *analysis.indicators.md5])


def _relative_pdf_link(output_file: Path, local_file_path: str) -> str:
    try:
        return os.path.relpath(local_file_path, output_file.parent)
    except ValueError:
        return local_file_path
